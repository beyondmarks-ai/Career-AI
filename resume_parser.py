"""Safe text extraction for supported resume formats."""

from pathlib import Path

from docx import Document
from PyPDF2 import PdfReader

MAX_FILE_SIZE = 8 * 1024 * 1024
MAX_RESUME_CHARS = 30_000


class ResumeParseError(ValueError):
    """A user-facing resume parsing error."""


def extract_resume_text(uploaded_file) -> str:
    """Extract text from a Streamlit UploadedFile (PDF or DOCX)."""
    file_name = getattr(uploaded_file, "name", "")
    extension = Path(file_name).suffix.lower()
    size = getattr(uploaded_file, "size", None)

    if extension not in {".pdf", ".docx"}:
        raise ResumeParseError("Please upload a PDF or DOCX resume.")
    if size is not None and size > MAX_FILE_SIZE:
        raise ResumeParseError("The resume is larger than 8 MB.")

    uploaded_file.seek(0)
    if extension == ".pdf":
        reader = PdfReader(uploaded_file)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ResumeParseError("Password-protected PDFs are not supported.") from exc
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    else:
        document = Document(uploaded_file)
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_cells = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
        text = "\n".join(paragraphs + table_cells)

    cleaned = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if len(cleaned) < 80:
        raise ResumeParseError(
            "Very little text could be extracted. Please use a text-based PDF or DOCX."
        )
    return cleaned[:MAX_RESUME_CHARS]
