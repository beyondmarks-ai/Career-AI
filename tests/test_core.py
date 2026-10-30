import io
import os
import unittest
from unittest.mock import Mock, patch

from azure.core.credentials import AzureKeyCredential
from docx import Document
from pydantic import ValidationError

from analysis_service import ResumeAnalysis
from learning_resources import learning_resources
from job_service import search_jobs_india
from rag_service import _search_credential, chunk_resume, retrieve_chat_context
from resume_parser import ResumeParseError, extract_resume_text


class NamedBytesIO(io.BytesIO):
    def __init__(self, data: bytes, name: str):
        super().__init__(data)
        self.name = name
        self.size = len(data)


class ResumeParserTests(unittest.TestCase):
    def test_extracts_docx_paragraphs_and_tables(self):
        document = Document()
        document.add_heading('Jordan Candidate', level=1)
        document.add_paragraph(
            'Data analyst with five years of experience building measurable dashboards.'
        )
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = 'Skills: Python, SQL, Power BI, stakeholder communication'
        buffer = io.BytesIO()
        document.save(buffer)
        upload = NamedBytesIO(buffer.getvalue(), 'resume.docx')

        text = extract_resume_text(upload)

        self.assertIn('Jordan Candidate', text)
        self.assertIn('Python, SQL', text)

    def test_rejects_unsupported_format(self):
        upload = NamedBytesIO(b'plain resume content', 'resume.txt')
        with self.assertRaises(ResumeParseError):
            extract_resume_text(upload)


class AnalysisSchemaTests(unittest.TestCase):
    def test_rejects_score_above_100(self):
        with self.assertRaises(ValidationError):
            ResumeAnalysis.model_validate(
                {
                    'overall_score': 101,
                    'ats_score': 80,
                    'role_match_score': 75,
                    'summary': 'Summary',
                    'strengths': ['One', 'Two'],
                    'improvements': ['One', 'Two', 'Three'],
                    'missing_keywords': [],
                    'skill_recommendations': [
                        {'name': str(i), 'priority': 'High', 'reason': 'Why', 'next_step': 'Do'}
                        for i in range(3)
                    ],
                    'action_plan': ['One', 'Two', 'Three'],
                }
            )


class ResumeChunkingTests(unittest.TestCase):
    def test_chunks_long_resume_with_overlap(self):
        resume = "\n".join(
            f"Project {index}: Built an Azure analytics workflow with measurable impact."
            for index in range(80)
        )

        chunks = chunk_resume(resume, max_chars=500, overlap_chars=60)

        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(chunk) <= 600 for chunk in chunks))
        self.assertIn("Azure analytics", " ".join(chunks))

    def test_chat_retrieval_selects_question_relevant_resume_evidence(self):
        resume = "\n".join(
            ["Jordan Candidate", "Education: BSc Economics"]
            + [f"General experience item {index} with routine responsibilities." for index in range(60)]
            + ["Acme Corp: Reduced reporting time by 35% using SQL and Power BI."]
        )

        context = retrieve_chat_context(
            resume,
            "What measurable data analytics achievement do I have?",
            "Data Analyst",
            max_chars=4000,
        )

        self.assertIn("Reduced reporting time by 35%", context)
        self.assertIn("Jordan Candidate", context)

    def test_uses_search_api_key_when_configured(self):
        with patch.dict(os.environ, {"AZURE_SEARCH_API_KEY": "test-search-key"}):
            credential = _search_credential()

        self.assertIsInstance(credential, AzureKeyCredential)
        self.assertEqual(credential.key, "test-search-key")


class LearningResourceTests(unittest.TestCase):
    def test_builds_safe_encoded_provider_links(self):
        resources = learning_resources("Power BI & DAX")

        self.assertEqual(len(resources), 4)
        self.assertTrue(all(resource.url.startswith("https://") for resource in resources))
        self.assertTrue(all(" " not in resource.url for resource in resources))
        self.assertIn("Power+BI+%26+DAX", resources[0].url)


class JoobleJobTests(unittest.TestCase):
    @patch("job_service.load_dotenv")
    @patch("job_service.requests.post")
    def test_searches_india_and_sanitizes_results(self, post, _load_dotenv):
        response = Mock(status_code=200)
        response.json.return_value = {
            "totalCount": 3,
            "jobs": [
                {
                    "id": 1,
                    "title": "Data <b>Analyst</b>",
                    "company": "Example Ltd",
                    "location": "Bengaluru, Karnataka",
                    "snippet": "<strong>Build</strong> dashboards &amp; reports.",
                    "type": "Full-time",
                    "link": "https://in.jooble.org/job/1",
                },
                {
                    "id": 2,
                    "title": "Unsafe job",
                    "link": "javascript:alert(1)",
                },
                {
                    "id": 3,
                    "title": "Duplicate",
                    "link": "https://in.jooble.org/job/1",
                },
            ],
        }
        post.return_value = response

        with patch.dict(os.environ, {"JOOBLE_API_KEY": "test-key"}):
            result = search_jobs_india("Data Analyst", 12)

        request = post.call_args
        self.assertEqual(request.kwargs["json"]["location"], "India")
        self.assertEqual(request.kwargs["json"]["keywords"], "Data Analyst")
        self.assertEqual(len(result.jobs), 1)
        self.assertEqual(result.jobs[0].title, "Data Analyst")
        self.assertEqual(result.jobs[0].snippet, "Build dashboards & reports.")


if __name__ == '__main__':
    unittest.main()
